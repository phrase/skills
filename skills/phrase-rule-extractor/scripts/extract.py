#!/usr/bin/env python3
"""Read translatable text segments from a file of almost any format.

Usage (standalone, to check what a file yields):
  python3 extract.py <file> [--side src|tgt] [--limit 20]

Used as a module by align.py and pair_files.py:
  from extract import load
  keyed, segments = load(path, side="tgt")   # segments: list of (key, text)

keyed=True means each segment has a stable key that is the same in the source, candidate, and
reference files (CMS JSON field paths, XLIFF unit ids, PO msgids, spreadsheet cells). Then
align.py pairs segments by key. keyed=False means the file is running text (PDF, DOCX, PPTX,
HTML, Markdown, plain text). Then align.py pairs segments by order and content.

Supported formats:
  .json                 CMS exports (strings under "value"/"title"), or any JSON with string leaves
  .xliff .xlf .mxliff .sdlxliff   <source> for side=src, <target> for side=tgt
  .po .pot              msgid for side=src, msgstr for side=tgt
  .csv .tsv .xlsx .xlsm one segment per non-empty text cell
  .docx                 body paragraphs and table cells in document order
  .doc .odt .rtf .pages converted to .docx with LibreOffice first
  .pptx                 slide text, tables, and speaker notes
  .ppt .odp             converted to .pptx with LibreOffice first
  .pdf                  text layer, rebuilt into paragraphs (scanned PDFs need OCR first)
  .html .htm .xhtml     block elements (p, li, td, th, h1-h6, title, figcaption)
  .md .markdown .txt    one segment per non-empty line
  .srt .vtt             subtitle text without numbers and timecodes
  anything else         read as UTF-8 text, one segment per non-empty line
"""
import argparse
import csv
import html
import json
import os
import re
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET

TEXT_KEYS = {"value", "title"}
SKIP_KEYS = {"uri", "id", "uuid", "href", "url", "src", "langcode", "format", "type"}
BLOCK = re.compile(r"</td>|</th>|</li>|</sub>|</p>|</h[1-6]>|<br\s*/?>|\r?\n")


class ExtractError(Exception):
    pass


def clean(part):
    """Remove markup and footnote markers, unescape entities, normalize spaces."""
    part = re.sub(r"<sup>.*?</sup>", "", part, flags=re.S)
    part = re.sub(r"<[^>]+>", "", part)
    part = html.unescape(html.unescape(part)).replace("\xa0", " ")
    return re.sub(r"\s+", " ", part).strip()


def split(text):
    """Split one string into block segments (paragraphs, cells, line breaks)."""
    return [c for c in (clean(p) for p in BLOCK.split(text)) if c]


def looks_like_text(s):
    s = s.strip()
    if not s or re.fullmatch(r"https?://\S+|[\w.-]+@[\w.-]+|[\d\s.,:/-]+|#?[0-9a-fA-F]{6}", s):
        return False
    return bool(re.search(r"[^\W\d_]", s))


# ---------- keyed formats ----------

def _json(path):
    data = json.load(open(path, encoding="utf-8-sig"))

    def walk(node, p, out, only_text_keys):
        if isinstance(node, dict):
            for k, v in node.items():
                if k.startswith("_") or k in SKIP_KEYS:
                    continue
                q = f"{p}/{k}" if p else k
                if isinstance(v, str):
                    if (k in TEXT_KEYS or not only_text_keys) and looks_like_text(v):
                        out.append((q, v))
                else:
                    walk(v, q, out, only_text_keys)
        elif isinstance(node, list):
            for i, item in enumerate(node):
                if isinstance(item, str) and not only_text_keys and looks_like_text(item):
                    out.append((f"{p}/{i}", item))
                else:
                    walk(item, f"{p}/{i}", out, only_text_keys)
        return out

    segs = walk(data, "", [], True)
    if not segs:  # not a CMS export: take every text-like string leaf, e.g. i18n key files
        segs = walk(data, "", [], False)
    return True, segs


def _strip_ns(root):
    for el in root.iter():
        if isinstance(el.tag, str) and "}" in el.tag:
            el.tag = el.tag.split("}", 1)[1]
    return root


def _inner_text(el):
    return "".join(el.itertext()) if el is not None else ""


def _xliff(path, side):
    root = _strip_ns(ET.parse(path).getroot())
    out = []
    for i, unit in enumerate(u for u in root.iter() if u.tag in ("trans-unit", "unit")):
        key = unit.get("id") or str(i)
        if unit.tag == "unit":  # XLIFF 2: unit/segment/source|target
            for j, seg in enumerate(unit.iter("segment")):
                txt = _inner_text(seg.find("source" if side == "src" else "target"))
                if txt.strip():
                    out.append((f"{key}/{j}", txt))
        else:
            tag = "source" if side == "src" else "target"
            el = unit.find(tag)
            if el is None and side == "tgt":
                continue
            txt = _inner_text(el)
            if txt.strip():
                out.append((key, txt))
    return True, out


def _po(path, side):
    out, cur, field = [], {}, None
    def flush():
        if cur.get("msgid"):
            out.append((cur["msgid"], cur.get("msgid" if side == "src" else "msgstr", "")))
    for line in open(path, encoding="utf-8"):
        line = line.strip()
        m = re.match(r'(msgid|msgstr)(?:\[0\])?\s+"(.*)"$', line)
        if m:
            if m.group(1) == "msgid":
                flush()
                cur = {}
            field = m.group(1)
            cur[field] = bytes(m.group(2), "utf-8").decode("unicode_escape").encode("latin-1").decode("utf-8")
        elif line.startswith('"') and field:
            cur[field] += bytes(line[1:-1], "utf-8").decode("unicode_escape").encode("latin-1").decode("utf-8")
        elif not line:
            field = None
    flush()
    return True, [(k, v) for k, v in out if v.strip()]


def _table(path):
    out = []
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and looks_like_text(c.value):
                        out.append((f"{ws.title}!{c.coordinate}", c.value))
    else:
        delim = "\t" if ext == ".tsv" else None
        text = open(path, encoding="utf-8-sig", newline="").read()
        dialect = csv.Sniffer().sniff(text[:4096]) if delim is None else None
        reader = csv.reader(text.splitlines(), dialect) if dialect else csv.reader(text.splitlines(), delimiter=delim)
        for r, row in enumerate(reader, 1):
            for c, val in enumerate(row, 1):
                if looks_like_text(val):
                    out.append((f"R{r}C{c}", val))
    return True, out


# ---------- running-text formats ----------

def _convert(path, target_ext):
    tmp = tempfile.mkdtemp()
    try:
        subprocess.run(["soffice", "--headless", "--convert-to", target_ext.lstrip("."), "--outdir", tmp, path],
                       check=True, capture_output=True, timeout=180)
    except (OSError, subprocess.SubprocessError) as e:
        raise ExtractError(f"Could not convert {path} with LibreOffice: {e}")
    out = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + target_ext)
    if not os.path.exists(out):
        raise ExtractError(f"LibreOffice produced no {target_ext} file for {path}")
    return out


def _docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    out = []

    def paras_of(container):
        for child in container.element.body.iterchildren() if hasattr(container.element, "body") else container._element.iterchildren():
            tag = child.tag.split("}")[1]
            if tag == "p":
                yield Paragraph(child, container).text
            elif tag == "tbl":
                for row in Table(child, container).rows:
                    seen = set()
                    for cell in row.cells:
                        if id(cell._tc) in seen:  # merged cells repeat
                            continue
                        seen.add(id(cell._tc))
                        for p in cell.paragraphs:
                            yield p.text

    for t in paras_of(d):
        out.append(t)
    for sec in d.sections:  # headers and footers last
        for part in (sec.header, sec.footer):
            for p in part.paragraphs:
                out.append(p.text)
    return False, [(f"p/{i}", t) for i, t in enumerate(x for x in out if x.strip())]


def _pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []

    def shape_texts(shape):
        if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
            for s in shape.shapes:
                yield from shape_texts(s)
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                yield "".join(r.text for r in p.runs)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text

    for n, slide in enumerate(prs.slides, 1):
        for shape in sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
            for t in shape_texts(shape):
                if t.strip():
                    out.append((f"slide{n}", t))
        if slide.has_notes_slide:
            for p in slide.notes_slide.notes_text_frame.paragraphs:
                if p.text.strip():
                    out.append((f"slide{n}/notes", p.text))
    return False, [(f"{k}/{i}", t) for i, (k, t) in enumerate(out)]


def _join_lines(lines):
    joined = lines[0]
    for l in lines[1:]:
        if joined.endswith("-") and l[:1].islower() and not joined[-2:-1].isdigit():
            joined = joined[:-1] + l          # word broken at the line end
        elif joined.endswith("-"):
            joined += l                       # compound such as Fußball-Schuhkollektion
        else:
            joined += " " + l
    return joined


PAGE_NO = re.compile(r"(page|seite|p\.)?\s*\d+(\s*(/|of|von)\s*\d+)?", re.I)


def _pdf(path):
    """Rebuild paragraphs from line positions: a larger vertical gap, or a short line that ends a
    sentence, starts a new paragraph. Repeated header and footer lines are removed."""
    try:
        import pdfplumber
    except ImportError:
        pdfplumber = None
    pages = []
    if pdfplumber:
        try:
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    words = page.extract_words(keep_blank_chars=False, use_text_flow=True)
                    lines = []
                    for w in words:
                        if lines and abs(w["top"] - lines[-1]["top"]) < 2.5:
                            lines[-1]["text"] += " " + w["text"]
                            lines[-1]["x1"] = max(lines[-1]["x1"], w["x1"])
                            lines[-1]["bottom"] = max(lines[-1]["bottom"], w["bottom"])
                        else:
                            lines.append({"text": w["text"], "top": w["top"], "bottom": w["bottom"],
                                          "x0": w["x0"], "x1": w["x1"]})
                    for l in lines:
                        l["edge"] = l["top"] < 0.1 * page.height or l["bottom"] > 0.9 * page.height
                    pages.append(lines)
        except Exception:
            pages = []
    if not pages:
        try:
            text = subprocess.run(["pdftotext", "-enc", "UTF-8", path, "-"], capture_output=True,
                                  check=True, timeout=180).stdout.decode("utf-8", "replace")
        except (OSError, subprocess.SubprocessError) as e:
            raise ExtractError(f"Could not read {path}: {e}")
        blocks = [[l.strip() for l in b.splitlines() if l.strip()] for b in re.split(r"\n\s*\n|\f", text)]
        paras = [_join_lines(b) for b in blocks if b]
    else:
        if sum(len(l["text"]) for p in pages for l in p) < 20:
            raise ExtractError(f"{path} has no text layer. It is probably scanned. Run OCR first, "
                               "for example: ocrmypdf in.pdf out.pdf")
        # header and footer lines that repeat on most pages
        from collections import Counter
        rep = Counter(l["text"] for p in pages for l in {x["text"]: x for x in p if x["edge"]}.values())
        repeated = {t for t, n in rep.items() if len(pages) > 2 and n >= 0.6 * len(pages)}
        paras, page_paras = [], []
        for lines in pages:
            page_paras.append(len(paras))
            lines = [l for l in lines if l["text"] not in repeated and not PAGE_NO.fullmatch(l["text"].strip())]
            if not lines:
                continue
            gaps = sorted(b["top"] - a["bottom"] for a, b in zip(lines, lines[1:]) if b["top"] > a["top"])
            normal = gaps[len(gaps) // 2] if gaps else 0
            right = max(l["x1"] for l in lines)
            cur = [lines[0]["text"]]
            for a, b in zip(lines, lines[1:]):
                gap = b["top"] - a["bottom"]
                short_end = a["x1"] < right - 0.12 * (right - min(l["x0"] for l in lines)) \
                    and re.search(r"[.!?:\u201c\u201d\"]$", a["text"])
                if gap > normal + 3 or b["top"] < a["top"] or short_end:
                    paras.append(_join_lines(cur))
                    cur = []
                cur.append(b["text"])
            paras.append(_join_lines(cur))
        page_starts = set(page_paras)
        # a paragraph that continues on the next page: no sentence end, next starts in lower case
        merged = []
        for i, p in enumerate(paras):
            if merged and i in page_starts and not re.search(r"[.!?:\u201c\u201d\"]$", merged[-1]) and (p[:1].islower() or len(merged[-1]) > 60):
                merged[-1] = _join_lines([merged[-1], p])
            else:
                merged.append(p)
        paras = merged
    paras = [p for p in paras if p.strip()]
    if len(re.sub(r"\s", "", " ".join(paras))) < 20:
        raise ExtractError(f"{path} has no text layer. It is probably scanned. Run OCR first, "
                           "for example: ocrmypdf in.pdf out.pdf")
    return False, [(f"p/{i}", t) for i, t in enumerate(paras)]


def _html(path):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(open(path, encoding="utf-8", errors="replace").read(), "lxml")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    for s in soup.find_all("sup"):
        s.decompose()
    blocks = ["title", "h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th",
              "figcaption", "caption", "dt", "dd", "blockquote"]
    out = []
    for el in soup.find_all(blocks + ["a", "button"]):
        if el.name in ("a", "button"):
            if el.find_parent(blocks):
                continue  # already part of its block
        elif el.find(blocks):
            continue  # take the innermost block only
        t = clean(el.get_text(" "))
        if t:
            out.append(t)
    for el in soup.find_all(attrs={"alt": True}):
        if looks_like_text(el["alt"]):
            out.append(el["alt"])
    return False, [(f"p/{i}", t) for i, t in enumerate(out)]


def _subtitles(path):
    out = []
    for block in re.split(r"\n\s*\n", open(path, encoding="utf-8-sig").read()):
        lines = [l for l in block.splitlines() if l.strip() and not re.fullmatch(r"\d+", l.strip())
                 and "-->" not in l and l.strip() != "WEBVTT"]
        if lines:
            out.append(" ".join(lines))
    return False, [(f"cue/{i}", t) for i, t in enumerate(out)]


def _lines(path):
    try:
        text = open(path, encoding="utf-8-sig").read()
    except UnicodeDecodeError:
        text = open(path, encoding="cp1252", errors="replace").read()
    lines = [l for l in text.splitlines() if l.strip()]
    if os.path.splitext(path)[1].lower() in (".md", ".markdown"):
        lines = [re.sub(r"^\s{0,3}(#{1,6}\s+|[-*+]\s+|\d+[.)]\s+|>\s*)", "", l) for l in lines]
        lines = [re.sub(r"(\*\*|__|`)", "", l) for l in lines if not re.fullmatch(r"\s*[-|:= ]{3,}\s*", l)]
    return False, [(f"line/{i}", l) for i, l in enumerate(lines)]


def load(path, side="tgt"):
    """Return (keyed, [(key, text), ...]). side is 'src' or 'tgt' for bilingual formats."""
    if not os.path.exists(path):
        raise ExtractError(f"File not found: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        try:
            return _json(path)
        except json.JSONDecodeError:
            return _lines(path)
    if ext in (".xliff", ".xlf", ".mxliff", ".sdlxliff", ".mqxliff"):
        return _xliff(path, side)
    if ext in (".po", ".pot"):
        return _po(path, side)
    if ext in (".csv", ".tsv", ".xlsx", ".xlsm"):
        return _table(path)
    if ext == ".xls":
        return _table(_convert(path, ".xlsx"))
    if ext == ".docx":
        return _docx(path)
    if ext in (".doc", ".odt", ".rtf", ".pages", ".wpd"):
        return _docx(_convert(path, ".docx"))
    if ext == ".pptx":
        return _pptx(path)
    if ext in (".ppt", ".odp", ".key"):
        return _pptx(_convert(path, ".pptx"))
    if ext == ".pdf":
        return _pdf(path)
    if ext in (".html", ".htm", ".xhtml"):
        return _html(path)
    if ext in (".srt", ".vtt"):
        return _subtitles(path)
    return _lines(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--side", default="tgt", choices=["src", "tgt"])
    ap.add_argument("--limit", type=int, default=20)
    a = ap.parse_args()
    try:
        keyed, segs = load(a.file, a.side)
    except ExtractError as e:
        sys.exit(str(e))
    parts = [(k, p) for k, t in segs for p in split(t)]
    print(f"{a.file}: {'keyed' if keyed else 'running text'}, {len(segs)} strings, {len(parts)} segments")
    for k, p in parts[:a.limit]:
        print(f"  [{k}] {p[:100]}")


if __name__ == "__main__":
    main()
